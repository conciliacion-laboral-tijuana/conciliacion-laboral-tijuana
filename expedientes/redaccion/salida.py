"""Salida de la demanda por capítulos: HTML y DOCX.

El DOCX se construye DIRECTAMENTE con python-docx (no vía `html_a_docx`) por dos
razones:

1. `html_a_docx` está pensado para el HTML de Quill del editor (una etiqueta por
   línea) y pierde el contenido cuando las etiquetas comparten línea.
2. Un escrito judicial necesita control tipográfico real: texto justificado,
   interlineado, encabezados de capítulo y tablas.

El HTML se usa para la vista previa en el navegador.
"""

from html import escape as _escape

from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
from docx.shared import Cm, Pt, RGBColor

from .demanda import capitulos, encabezado_capitulo

# ─── Tipografía del escrito (márgenes y cuerpo) ──────────────────────────────
FUENTE = 'Arial'
TAMANO_CUERPO = 11
TAMANO_ENCABEZADO = 12
TAMANO_TITULO = 14
INTERLINEADO = 1.5


def _e(valor) -> str:
    return _escape(str(valor or ''))


def _objeto_html(objeto) -> str:
    """Renderiza el 'objeto' de una prueba (str o lista numerada) como HTML."""
    if isinstance(objeto, (list, tuple)):
        items = ''.join(f'<li>{_e(o)}</li>' for o in objeto)
        return (f'<p>Lo anterior con el fin de acreditar el objeto:</p>'
                f'<ol>{items}</ol>')
    if objeto:
        return f'<p>{_e(objeto)}</p>'
    return ''


# ══════════════════════════════════════════════════════════════════════════
# HTML (vista previa)
# ══════════════════════════════════════════════════════════════════════════

# Orden de los bloques en el escrito. La pantalla del documento usa el mismo
# orden para el índice de capítulos, así que vive aquí y no en la plantilla.
ORDEN_HTML = ('encabezado', 'proemio', 'demandados', 'prestaciones', 'jornada',
              'hechos', 'pruebas', 'derecho', 'cierre')


def capitulos_html(datos: dict, calculo: dict,
                   tipo_despido: str = 'injustificado') -> dict:
    """HTML de la demanda SEPARADO por capítulos: {capítulo: html}.

    La pantalla del documento necesita cada capítulo por separado para poder
    señalar en cuál falta algo; `generar_demanda_capitulos_html` los vuelve a
    unir en el mismo orden para la descarga y las pruebas.
    """
    cap = capitulos(datos, calculo, tipo_despido)
    car = cap['caratula']
    bloques = {}

    bloques['encabezado'] = '\n'.join([
        '<p><strong>INICIO.</strong></p>',
        f'<p><strong>ACTOR.</strong> {_e(car["actor"])}</p>',
        f'<p><strong>ACCIÓN.</strong> {_e(car["accion"])}</p>',
        f'<h2>{_e(car["tribunal"])}</h2>',
        '<p><strong>P R E S E N T E.-</strong></p>',
    ])

    bloques['proemio'] = '\n'.join(
        f'<p>{_e(parrafo)}</p>' for parrafo in cap['proemio'].split('\n\n')
    )

    demandados = []
    if cap['demandados']:
        demandados.append('<p>Que mediante el presente escrito y con fundamento en, '
                          'vengo a demandar a:</p>')
        romanos = ['I', 'II', 'III', 'IV', 'V']
        for i, numeral in enumerate(cap['demandados']):
            etiqueta = romanos[i] if i < len(romanos) else str(i + 1)
            demandados.append(f'<p><strong>{etiqueta}.-</strong> {_e(numeral)}</p>')
    demandados.append(f'<p>{_e(cap["introduccion"])}</p>')
    bloques['demandados'] = '\n'.join(demandados)

    prestaciones = []
    if cap['prestaciones']:
        prestaciones.append(f'<h3>{encabezado_capitulo("PRESTACIONES")}</h3>')
        for i, renglon in enumerate(cap['prestaciones'], start=1):
            prestaciones.append(f'<p>{i}. {_e(renglon)}</p>')
    bloques['prestaciones'] = '\n'.join(prestaciones)

    jornada = []
    if cap['jornada']['filas']:
        jornada.append(f'<h3>{encabezado_capitulo("JORNADA SEMANAL")}</h3>')
        if cap['jornada']['resumen']:
            jornada.append(f'<p>{_e(cap["jornada"]["resumen"])}</p>')
        enc = cap['jornada']['encabezados']
        cabecera = ''.join(f'<th>{_e(h)}</th>' for h in enc)
        filas = ''
        for fila in cap['jornada']['filas']:
            filas += (
                '<tr>'
                f'<td>{_e(fila["jornada_semanal"])}</td>'
                f'<td>{_e(fila["jornada_ordinaria"])}</td>'
                f'<td>{_e(fila["jornada_extra"])}</td>'
                f'<td>{_e(fila["horas"])}</td>'
                '</tr>'
            )
        jornada.append(f'<table><thead><tr>{cabecera}</tr></thead>'
                       f'<tbody>{filas}</tbody></table>')
    bloques['jornada'] = '\n'.join(jornada)

    hechos = []
    if cap['hechos']:
        hechos.append(f'<h3>{encabezado_capitulo("HECHOS")}</h3>')
        for hecho in cap['hechos']:
            hechos.append(f'<p>{_e(hecho)}</p>')
    bloques['hechos'] = '\n'.join(hechos)

    pruebas = []
    if cap['pruebas']:
        pruebas.append('<h3>PRUEBAS</h3>')
        for i, prueba in enumerate(cap['pruebas'], start=1):
            pruebas.append(
                f'<p><strong>{i}.- {_e(prueba["titulo"])}.</strong> '
                f'{_e(prueba["texto"])}</p>'
            )
            objeto = _objeto_html(prueba.get('objeto'))
            if objeto:
                pruebas.append(objeto)
    bloques['pruebas'] = '\n'.join(pruebas)

    derecho = []
    if cap['derecho']:
        derecho.append(f'<h3>{encabezado_capitulo("DERECHO")}</h3>')
        derecho.append(f'<p>{_e(cap["derecho"])}</p>')
    bloques['derecho'] = '\n'.join(derecho)

    bloques['cierre'] = '\n'.join([
        '<p>&nbsp;</p>',
        '<p><strong>PROTESTO LO NECESARIO</strong></p>',
        f'<p>{_e(car["actor"])}</p>',
    ])

    return bloques


def generar_demanda_capitulos_html(datos: dict, calculo: dict,
                                   tipo_despido: str = 'injustificado') -> str:
    """HTML de la demanda completa con la estructura de capítulos."""
    bloques = capitulos_html(datos, calculo, tipo_despido)
    return '\n'.join(bloques[clave] for clave in ORDEN_HTML if bloques.get(clave))


# ══════════════════════════════════════════════════════════════════════════
# DOCX (documento para presentar)
# ══════════════════════════════════════════════════════════════════════════

def _configurar(doc) -> None:
    """Márgenes y estilo base del escrito."""

    seccion = doc.sections[0]
    seccion.top_margin = Cm(2.5)
    seccion.bottom_margin = Cm(2.5)
    seccion.left_margin = Cm(3)
    seccion.right_margin = Cm(2.5)

    estilo = doc.styles['Normal']
    estilo.font.name = FUENTE
    estilo.font.size = Pt(TAMANO_CUERPO)
    estilo.paragraph_format.line_spacing = INTERLINEADO
    estilo.paragraph_format.space_after = Pt(6)
    estilo.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY


def _parrafo(doc, texto: str, negrita: bool = False, centrado: bool = False,
             justificado: bool = True, tamano: int | None = None,
             espacio_antes: int = 0):
    """Agrega un párrafo con el formato del escrito."""

    p = doc.add_paragraph()
    run = p.add_run(texto)
    run.bold = negrita
    if tamano:
        run.font.size = Pt(tamano)
    if centrado:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    elif justificado:
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    else:
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.line_spacing = INTERLINEADO
    if espacio_antes:
        p.paragraph_format.space_before = Pt(espacio_antes)
    return p


def _encabezado(doc, texto: str):
    """Encabezado de capítulo: negritas, centrado, con aire arriba."""
    return _parrafo(doc, texto, negrita=True, centrado=True,
                    justificado=False, tamano=TAMANO_ENCABEZADO,
                    espacio_antes=12)


def _parrafo_mixto(doc, partes: list):
    """Párrafo con fragmentos (texto, negrita) en una sola línea."""

    p = doc.add_paragraph()
    for texto, negrita in partes:
        run = p.add_run(texto)
        run.bold = negrita
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.line_spacing = INTERLINEADO
    p.paragraph_format.space_after = Pt(6)
    return p


def _lista_numerada(doc, items: list):
    """Lista numerada simple (1.-, 2.- …)."""

    for i, item in enumerate(items, start=1):
        p = doc.add_paragraph()
        run = p.add_run(f'{i}.- ')
        run.bold = True
        p.add_run(str(item))
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.line_spacing = INTERLINEADO
        p.paragraph_format.left_indent = None


def _tabla(doc, encabezados: list, filas: list):
    """Tabla con encabezado sombreado (el cuadro de jornada semanal)."""

    tabla = doc.add_table(rows=1, cols=len(encabezados))
    tabla.style = 'Table Grid'

    for i, texto in enumerate(encabezados):
        celda = tabla.rows[0].cells[i]
        celda.text = ''
        p = celda.paragraphs[0]
        run = p.add_run(str(texto))
        run.bold = True
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        celda._tc.get_or_add_tcPr().append(
            parse_xml(f'<w:shd {nsdecls("w")} w:fill="1F2937"/>')
        )

    for fila in filas:
        celdas = tabla.add_row().cells
        for i, valor in enumerate(fila):
            if i >= len(celdas):
                break
            celdas[i].text = ''
            p = celdas[i].paragraphs[0]
            run = p.add_run(str(valor))
            run.font.size = Pt(9)
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.line_spacing = 1.0

    doc.add_paragraph()


def generar_demanda_capitulos_docx(expediente, datos: dict | None = None,
                                   calculo: dict | None = None,
                                   tipo_despido: str | None = None):
    """Documento Word con la demanda completa por capítulos."""
    from docx import Document

    from .adaptador import datos_de_expediente

    tipo = tipo_despido or expediente.tipo_despido or 'injustificado'
    if datos is None:
        datos = datos_de_expediente(expediente)
    if calculo is None:
        from ..demanda_generator import calculo_para_demanda
        calculo = calculo_para_demanda(expediente, tipo)

    cap = capitulos(datos, calculo, tipo)
    car = cap['caratula']

    doc = Document()
    _configurar(doc)

    # ─── Carátula ─────────────────────────────────────────────────────────
    _parrafo(doc, 'INICIO.', negrita=True, justificado=False)
    _parrafo_mixto(doc, [('ACTOR. – ', True), (car['actor'], False)])
    _parrafo_mixto(doc, [('ACCIÓN. - ', True), (car['accion'], False)])
    _parrafo(doc, '', justificado=False)
    _parrafo(doc, car['tribunal'], negrita=True, centrado=True, justificado=False,
             tamano=TAMANO_TITULO)
    _parrafo(doc, 'P R E S E N T E.-', negrita=True, justificado=False,
             espacio_antes=6)

    # ─── Proemio ──────────────────────────────────────────────────────────
    for parrafo in cap['proemio'].split('\n\n'):
        _parrafo(doc, parrafo)

    # ─── Demandados ───────────────────────────────────────────────────────
    if cap['demandados']:
        _parrafo(doc, 'Que mediante el presente escrito y con fundamento en, '
                      'vengo a demandar a:')
        romanos = ['I', 'II', 'III', 'IV', 'V']
        for i, numeral in enumerate(cap['demandados']):
            etiqueta = romanos[i] if i < len(romanos) else str(i + 1)
            _parrafo_mixto(doc, [(f'{etiqueta}.- ', True), (numeral, False)])

    _parrafo(doc, cap['introduccion'])

    # ─── Capítulo de prestaciones ─────────────────────────────────────────
    if cap['prestaciones']:
        _encabezado(doc, encabezado_capitulo('PRESTACIONES'))
        _lista_numerada(doc, cap['prestaciones'])

    # ─── Anexo de jornada / horas extras ──────────────────────────────────
    if cap['jornada']['filas']:
        _encabezado(doc, encabezado_capitulo('JORNADA SEMANAL'))
        if cap['jornada']['resumen']:
            _parrafo(doc, cap['jornada']['resumen'])
        _tabla(doc, cap['jornada']['encabezados'], [
            [f['jornada_semanal'], f['jornada_ordinaria'],
             f['jornada_extra'], f['horas']]
            for f in cap['jornada']['filas']
        ])

    # ─── Capítulo de hechos ───────────────────────────────────────────────
    if cap['hechos']:
        _encabezado(doc, encabezado_capitulo('HECHOS'))
        for hecho in cap['hechos']:
            _parrafo(doc, hecho)

    # ─── Pruebas ──────────────────────────────────────────────────────────
    if cap['pruebas']:
        _encabezado(doc, 'PRUEBAS')
        for i, prueba in enumerate(cap['pruebas'], start=1):
            _parrafo_mixto(doc, [
                (f'{i}.- {prueba["titulo"]}. ', True),
                (prueba['texto'], False),
            ])
            objeto = prueba.get('objeto')
            if isinstance(objeto, (list, tuple)):
                _parrafo(doc, 'Lo anterior con el fin de acreditar el objeto:')
                _lista_numerada(doc, list(objeto))
            elif objeto:
                _parrafo(doc, str(objeto))

    # ─── Capítulo de derecho ──────────────────────────────────────────────
    if cap['derecho']:
        _encabezado(doc, encabezado_capitulo('DERECHO'))
        _parrafo(doc, cap['derecho'])

    # ─── Cierre ───────────────────────────────────────────────────────────
    _parrafo(doc, '', justificado=False)
    _parrafo(doc, 'PROTESTO LO NECESARIO', negrita=True, centrado=True,
             justificado=False)
    _parrafo(doc, car['actor'], negrita=True, centrado=True, justificado=False)

    return doc


def guardar_demanda_capitulos(expediente, ruta, datos: dict | None = None,
                              calculo: dict | None = None,
                              tipo_despido: str | None = None) -> str:
    """Genera el DOCX y lo guarda en `ruta`. Devuelve la ruta."""
    doc = generar_demanda_capitulos_docx(expediente, datos, calculo, tipo_despido)
    doc.save(str(ruta))
    return str(ruta)
