from datetime import date, timedelta
from decimal import Decimal as D

from django.contrib.auth.models import User
from django.core.exceptions import ValidationError
from django.test import TestCase
from django.urls import reverse

from .models import Cliente, Expediente, HojaPrestaciones, SolicitudConciliacion
from .hoja_prestaciones import (calcular_hoja, ciclos_para, empaquetar, desempaquetar,
                               formularios, validar_formularios, resultado_vigente, huella,
                               requisitos_para_calcular)
from .demanda_generator import construir_hechos, calculo_para_demanda, generar_demanda_html, _filas_prestaciones
from .laboral_calculator import calcular_desde_expediente


class HojaPrestacionesTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.user = User.objects.create_user('abogada_hoja')
        cls.user.profile.rol = 'abogada'
        cls.user.profile.save()
        cls.cliente = Cliente.objects.create(nombre='Persona de prueba', curp='AAAA900101HBCBBB01',
            salario=D('18000'), fecha_ingreso=date(2024, 1, 1), fecha_salida=date(2026, 7, 1),
            imss_confirmado=True, tuvo_imss=True, puesto='Operador', empresa='Empresa ejemplo')
        cls.exp = Expediente.objects.create(cliente=cls.cliente, asesor=cls.user, tipo_despido='injustificado')

    def setUp(self):
        self.client.force_login(self.user)
        self.cliente.refresh_from_db()
        self.exp.refresh_from_db()
        self.exp.cliente = self.cliente

    def datos(self):
        rows = []
        for ciclo in ciclos_para(self.cliente):
            row = dict(anio=ciclo.numero, dias_derecho=D(ciclo.dias_correspondientes),
                       disfrutados=D('0'), salario_diario=D('600'), pagado=D('0'),
                       prima_pagada=D('0'), incluir=True, confirmado=True)
            if ciclo.numero == 1:
                row['disfrutados'] = D(ciclo.dias_correspondientes)
                row['prima_pagada'] = D(ciclo.dias_correspondientes) * 600 * D('.25')
            rows.append(row)
        return empaquetar({'fechas': [str(self.cliente.fecha_ingreso), str(self.cliente.fecha_salida)],
            'config': {'extras_estado': 'no', 'datos_confirmados': True, 'aguinaldo_pagado': D('0'),
                       'dias_aguinaldo': D('15'), 'prima_porcentaje': D('25'),
                       'salario_integrado': D('650'), 'salarios_devengados': D('0')},
            'semanas': [], 'vacaciones': rows})

    def payload(self, datos=None, accion='calcular'):
        data = desempaquetar(datos or self.datos())
        payload = {'periodos_fechas': f'{self.cliente.fecha_ingreso}|{self.cliente.fecha_salida}', 'accion': accion, 'hoja_calculos': '1', 'imss_present': '1',
                   'imss_confirmado': 'on', 'tuvo_imss': 'on', 'tipo_despido': 'injustificado',
                   'semanas-TOTAL_FORMS': str(len(data['semanas'])), 'semanas-INITIAL_FORMS': str(len(data['semanas'])),
                   'vacaciones-TOTAL_FORMS': str(len(data['vacaciones'])), 'vacaciones-INITIAL_FORMS': str(len(data['vacaciones']))}
        for key, value in data['config'].items():
            if value is not False:
                payload['prestaciones-' + key] = 'on' if value is True else str(value)
        for prefix in ('semanas', 'vacaciones'):
            for i, row in enumerate(data[prefix]):
                for key, value in row.items():
                    if value is not False:
                        payload[f'{prefix}-{i}-{key}'] = 'on' if value is True else str(value)
        hoja = HojaPrestaciones.objects.filter(expediente=self.exp).first()
        payload['revision_calculo'] = hoja.huella if hoja else ''
        return payload

    def semana(self, lunes, salario='600', horas=(3, 3, 3, 3, 0, 0, 0), pagado='0'):
        return dict(semana=lunes, salario_diario=D(salario), jornada='diurna', limite_dobles=D('9'),
                    pagado=D(pagado), **dict(zip(('lunes','martes','miercoles','jueves','viernes','sabado','domingo'), map(D, horas))))

    def test_horas_irregulares_no_se_promedian_entre_semanas_y_no_se_truncan(self):
        datos = desempaquetar(self.datos())
        datos['config']['extras_estado'] = 'si'
        datos['semanas'] = [self.semana(date(2026, 6, 8), pagado='25'),
                            self.semana(date(2026, 6, 15), salario='400', horas=(3,0,0,0,0,0,0))]
        result = calcular_hoja(self.exp, empaquetar(datos))
        self.assertEqual(result['horas_extras']['dobles'], D('12'))
        self.assertEqual(result['horas_extras']['triples'], D('3'))
        self.assertEqual(result['horas_extras']['monto'], D('2300.00'))
        datos['semanas'] = [self.semana(date(2026, 6, 8), horas=(20,0,0,0,0,0,0))]
        result = calcular_hoja(self.exp, empaquetar(datos))
        self.assertEqual(result['horas_extras']['triples'], D('11'))

    def test_vacaciones_disfrutadas_prima_pendiente_y_proporcional_sin_duplicar(self):
        datos = desempaquetar(self.datos())
        result = calcular_hoja(self.exp, empaquetar(datos))
        ciclo_actual = ciclos_para(self.cliente)[-1]
        proporcional = (D(ciclo_actual.dias_correspondientes) * ciclo_actual.fraccion * 600).quantize(D('.01'))
        self.assertEqual(result['vacaciones']['monto'], D('8400') + proporcional)
        self.assertEqual(result['vacaciones_vencidas']['monto'], D('0'))
        self.assertEqual(result['prima_vacacional']['monto'], D('2100') + (proporcional * D('.25')).quantize(D('.01')))
        datos['vacaciones'][0]['prima_pagada'] = D('0')
        otro = calcular_hoja(self.exp, empaquetar(datos))
        self.assertEqual(otro['vacaciones']['monto'], result['vacaciones']['monto'])
        self.assertEqual(otro['prima_vacacional']['monto'] - result['prima_vacacional']['monto'], D('1800'))
        total_tabla = sum(row[2] for row in _filas_prestaciones(result, self.exp))
        self.assertEqual(total_tabla, result['total'])

    def test_pago_excesivo_y_dias_excesivos_bloquean_calculo(self):
        datos = desempaquetar(self.datos())
        datos['vacaciones'][1]['disfrutados'] = D('99')
        with self.assertRaises(ValidationError):
            calcular_hoja(self.exp, empaquetar(datos))
        datos = desempaquetar(self.datos())
        datos['config']['aguinaldo_pagado'] = D('999999')
        with self.assertRaises(ValidationError):
            calcular_hoja(self.exp, empaquetar(datos))

    def test_imss_condicional_no_presume_subregistro_y_html_escapado(self):
        result = calcular_desde_expediente(self.exp)
        hechos = ' '.join(construir_hechos(self.exp, result))
        self.assertIn('dado de alta ante el IMSS', hechos)
        self.assertNotIn('salario inferior', hechos)
        self.cliente.imss_salario_inferior = True
        self.cliente.imss_documento = '<script>alert(1)</script>'
        self.assertIn('salario inferior', ' '.join(construir_hechos(self.exp, result)))
        html = generar_demanda_html(self.exp)
        self.assertNotIn('<script>alert(1)</script>', html)
        self.cliente.tuvo_imss = False
        self.assertNotIn('dado de alta ante el IMSS', ' '.join(construir_hechos(self.exp, result)))

    def test_flujo_calcular_aprobar_generar_y_cambios_invalidan(self):
        url = reverse('demanda_asistente', args=[self.exp.pk])
        response = self.client.post(url, self.payload())
        self.assertEqual(response.status_code, 200)
        hoja = HojaPrestaciones.objects.get(expediente=self.exp)
        self.assertTrue(hoja.resultado)
        self.assertContains(response, 'name="revision_calculo" value="' + hoja.huella + '"')
        self.assertIsNone(hoja.aprobado_en)
        self.assertFalse(calculo_para_demanda(self.exp)['success'])
        response = self.client.post(url, self.payload(accion='aprobar'))
        self.assertEqual(response.status_code, 200)
        hoja.refresh_from_db()
        self.assertEqual(hoja.aprobado_por, self.user, str((response.context['errores'], response.context['captura_prestaciones'].errors, response.context['vacaciones_formset'].errors, response.context['semanas_formset'].errors)))
        self.assertTrue(calculo_para_demanda(self.exp)['success'])
        response = self.client.post(url, self.payload(accion='finalizar'))
        self.assertRedirects(response, reverse('demanda_editor', args=[self.exp.pk]))
        self.cliente.salario = D('19000')
        self.cliente.save()
        self.assertIsNone(resultado_vigente(self.exp, aprobado=True))
        self.assertFalse(calculo_para_demanda(self.exp)['success'])
        response = self.client.post(reverse('demanda_descargar', args=[self.exp.pk]), {'contenido': '<p>Anterior</p>'})
        self.assertRedirects(response, url)

    def test_no_aprueba_resultado_que_no_vio_y_borrador_no_calcula(self):
        url = reverse('demanda_asistente', args=[self.exp.pk])
        self.client.post(url, self.payload())
        payload = self.payload(accion='aprobar')
        payload['revision_calculo'] = 'revision_anterior'
        self.client.post(url, payload)
        hoja = HojaPrestaciones.objects.get(expediente=self.exp)
        self.assertIsNone(hoja.aprobado_en)
        self.client.post(url, self.payload(accion='aprobar'))
        payload = self.payload()
        payload['guardar_borrador'] = '1'
        payload['prestaciones-aguinaldo_pagado'] = '100'
        response = self.client.post(reverse('demanda_vista_previa', args=[self.exp.pk]), payload)
        self.assertEqual(response.status_code, 200)
        self.assertTrue(response.json()['calculo_pendiente'])
        hoja.refresh_from_db()
        self.assertIsNone(hoja.aprobado_en)

    def test_semanas_duplicadas_y_fuera_de_relacion_rechazadas(self):
        data = desempaquetar(self.datos())
        data['config']['extras_estado'] = 'si'
        semana = self.semana(date(2026, 6, 8))
        data['semanas'] = [semana, semana]
        forms = formularios(self.exp, self.payload(empaquetar(data)))
        self.assertIsNone(validar_formularios(self.exp, forms, final=True))
        data['semanas'] = [self.semana(date(2023, 1, 2))]
        forms = formularios(self.exp, self.payload(empaquetar(data)))
        self.assertIsNone(validar_formularios(self.exp, forms, final=True))

    def test_no_calcula_sin_confirmar_imss_o_periodos(self):
        self.cliente.imss_confirmado = False
        with self.assertRaises(ValidationError):
            calcular_hoja(self.exp, self.datos())
        self.cliente.imss_confirmado = True
        data = desempaquetar(self.datos())
        data['vacaciones'][0]['confirmado'] = False
        with self.assertRaises(ValidationError):
            calcular_hoja(self.exp, empaquetar(data))

    def test_solicitud_semanal_no_convierte_salario_mensual_en_semanal(self):
        SolicitudConciliacion.objects.create(expediente=self.exp, periodo_pago='semanal')
        result = calcular_desde_expediente(self.exp)
        self.assertEqual(result['salario_diario'], D('600'))

    def test_actualizar_fecha_reconstruye_periodos_y_exige_revisarlos(self):
        url = reverse('demanda_asistente', args=[self.exp.pk])
        self.client.post(url, self.payload())
        self.cliente.fecha_ingreso = date(2023, 1, 1)
        self.cliente.save()
        response = self.client.get(url)
        self.assertEqual(len(response.context['vacaciones_formset'].forms), 4)
        self.assertFalse(response.context['calculo_aprobado'])
        self.assertFalse(response.context['vacaciones_formset'].forms[0]['confirmado'].value())

    def test_docx_y_html_usan_desglose_aprobado_y_mismos_importes(self):
        from .demanda_generator import generar_demanda_word
        result = calcular_hoja(self.exp, self.datos())
        from django.utils import timezone
        HojaPrestaciones.objects.create(expediente=self.exp, datos=self.datos(),
            resultado=empaquetar(result), huella=huella(self.exp, self.datos()),
            aprobado_por=self.user, aprobado_en=timezone.now())
        html = generar_demanda_html(self.exp)
        doc = generar_demanda_word(self.exp)
        texto = ' '.join(p.text for p in doc.paragraphs)
        self.assertIn('Desglose de prestaciones por periodos', html)
        self.assertIn('Desglose de prestaciones por periodos', texto)
        self.assertIn('dado de alta ante el IMSS', texto)
        self.assertIn('2025-01-01 a 2025-12-31', html)
        self.assertIn('2025-01-01 a 2025-12-31', texto)

    def test_salario_cambia_dentro_de_la_semana(self):
        data = desempaquetar(self.datos())
        data['config']['extras_estado'] = 'si'
        row = self.semana(date(2026, 6, 8), salario='400', horas=(9,3,0,0,0,0,0))
        row['salario_martes'] = D('600')
        data['semanas'] = [row]
        result = calcular_hoja(self.exp, empaquetar(data))
        self.assertEqual(result['horas_extras']['monto'], D('1575.00'))

    def test_fila_nueva_sigue_existiendo_despues_de_autoguardar(self):
        data = desempaquetar(self.datos())
        data['config']['extras_estado'] = 'si'
        data['semanas'] = [self.semana(date(2026, 6, 8))]
        url = reverse('demanda_vista_previa', args=[self.exp.pk])
        payload = self.payload(empaquetar(data))
        payload['semanas-INITIAL_FORMS'] = '0'
        payload['guardar_borrador'] = '1'
        self.client.post(url, payload)
        response = self.client.post(reverse('demanda_asistente', args=[self.exp.pk]), payload)
        self.assertEqual(response.status_code, 200)
        hoja = HojaPrestaciones.objects.get(expediente=self.exp)
        self.assertEqual(len(hoja.datos['semanas']), 1)
        self.assertTrue(hoja.resultado)

    def test_corregir_nombre_no_obliga_a_recalcular_importes(self):
        from django.utils import timezone
        data = self.datos()
        HojaPrestaciones.objects.create(expediente=self.exp, datos=data,
            resultado=empaquetar(calcular_hoja(self.exp, data)), huella=huella(self.exp, data),
            aprobado_por=self.user, aprobado_en=timezone.now())
        self.cliente.nombre = 'Nombre corregido'
        self.cliente.save()
        self.assertIsNotNone(resultado_vigente(self.exp, aprobado=True))

    def test_no_usa_monto_reclamado_anterior_sin_calculo_aprobado(self):
        self.exp.monto_reclamado = D('99999')
        self.exp.save()
        HojaPrestaciones.objects.create(expediente=self.exp)
        html = generar_demanda_html(self.exp)
        self.assertNotIn('$99,999.00', html)
        self.assertFalse(calculo_para_demanda(self.exp)['success'])

    def test_horarios_coherentes_se_guardan_y_generan_cuadro_aprobado(self):
        datos = desempaquetar(self.datos())
        datos['config']['extras_estado'] = 'si'
        row = self.semana(date(2026, 6, 29), horas=(0,0, D('1.6'),0,0,0,0), pagado='20')
        row.update(entrada_miercoles='07:00', ordinaria_miercoles='15:24', salida_miercoles='17:00')
        datos['semanas'] = [row]
        context = formularios(self.exp, self.payload(empaquetar(datos)))
        stored = validar_formularios(self.exp, context, final=True)
        self.assertIsNotNone(stored)
        result = calcular_hoja(self.exp, stored)
        self.assertEqual(result['horas_extras']['monto'], D('220'))
        self.assertEqual(result['horas_extras']['semanas'][0]['fin'], date(2026,7,1))
        from .demanda_generator import _cuadro_horas_extras_html, html_a_docx
        from docx import Document
        cuadro = _cuadro_horas_extras_html(result)
        self.assertIn('1 hora y 36 minutos', cuadro)
        self.assertIn('15:24 a 17:00', cuadro)
        self.assertIn('$220.00', cuadro)
        doc = Document()
        html_a_docx(cuadro, doc)
        self.assertEqual(len(doc.tables), 1)
        self.assertIn('15:24', doc.tables[0].cell(1,2).text)

    def test_horarios_incompletos_o_inconsistentes_no_se_aprueban(self):
        from .hoja_prestaciones import SemanaForm
        row = self.semana(date(2026,6,8), horas=(2,0,0,0,0,0,0))
        row.update(entrada_lunes='07:00', ordinaria_lunes='15:00', salida_lunes='17:00')
        self.assertTrue(SemanaForm(row).is_valid())
        row['salida_lunes'] = '18:00'
        self.assertFalse(SemanaForm(row).is_valid())
        row['salida_lunes'] = ''
        self.assertFalse(SemanaForm(row).is_valid())
        row.update(entrada_lunes='22:00', ordinaria_lunes='06:00', salida_lunes='08:00')
        self.assertTrue(SemanaForm(row).is_valid())


class RequisitosParaCalcularTests(HojaPrestacionesTests):
    """La pantalla enumera qué falta para calcular y enlaza a cada campo.

    El defecto que esto impide: «los importes aparecerán después de validar y
    calcular» sin decir qué faltaba ni dónde capturarlo, con el botón a más de
    cien campos de distancia del mensaje.
    """

    def _requisitos(self, post=None):
        return requisitos_para_calcular(self.exp, formularios(self.exp, post))

    def _por_clave(self, post=None):
        return {r['clave']: r for r in self._requisitos(post)}

    def test_un_caso_nuevo_enumera_lo_que_falta(self):
        """El fixture ya trae IMSS confirmado y fechas: el resto está pendiente."""
        pendientes = {clave for clave, r in self._por_clave().items()
                      if not r['cumplido']}

        self.assertEqual(pendientes, {'salario_integrado', 'datos_confirmados',
                                      'extras_estado', 'vacaciones'})
        self.assertTrue(self._por_clave()['fechas']['cumplido'])

    def test_cada_requisito_sabe_llevar_al_campo(self):
        """Sin destino, la lista sería un diagnóstico sin salida."""
        for requisito in self._requisitos():
            self.assertTrue(requisito['ancla'], requisito['clave'])
            self.assertTrue(requisito['bloque'], requisito['clave'])
            self.assertIn(requisito['espacio'], ('datos', 'estrategia'), requisito['clave'])

    def test_con_todo_capturado_no_queda_ninguno_pendiente(self):
        """El mismo POST que calcula tiene que dejar la lista en verde."""
        pendientes = [r for r in self._requisitos(self.payload()) if not r['cumplido']]

        self.assertEqual(pendientes, [])

    def test_sin_confirmar_el_imss_lo_reclama_y_lleva_a_su_bloque(self):
        self.cliente.imss_confirmado = False
        self.cliente.save(update_fields=['imss_confirmado'])

        requisito = self._por_clave()['imss_confirmado']

        self.assertFalse(requisito['cumplido'])
        self.assertEqual(requisito['ancla'], 'sec-imss')
        self.assertEqual(requisito['bloque'], 'imss')

    def test_sin_fechas_lo_reclama_en_el_espacio_de_datos(self):
        self.cliente.fecha_ingreso = None
        self.cliente.save(update_fields=['fecha_ingreso'])

        requisito = self._por_clave()['fechas']

        self.assertFalse(requisito['cumplido'])
        self.assertEqual(requisito['espacio'], 'datos')
        self.assertEqual(requisito['bloque'], 'laboral')

    def test_con_horas_extras_pide_detallar_las_semanas(self):
        datos = desempaquetar(self.datos())
        datos['config']['extras_estado'] = 'si'
        datos['semanas'] = []

        requisito = self._por_clave(self.payload(empaquetar(datos)))['semanas']

        self.assertFalse(requisito['cumplido'])

    def test_sin_horas_extras_no_se_pide_detallar_semanas(self):
        claves = set(self._por_clave().keys())
        self.assertNotIn('semanas', claves)

    def test_las_vacaciones_cuentan_los_periodos_confirmados(self):
        """No basta con que existan: cada aniversario se confirma por separado."""
        requisito = self._por_clave()['vacaciones']

        self.assertFalse(requisito['cumplido'])
        self.assertIn('0 de', requisito['ayuda'])
